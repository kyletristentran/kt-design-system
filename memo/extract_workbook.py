#!/usr/bin/env python3
"""
extract_workbook.py — turn an .xlsx into a structured JSON brief the model can reason about.

    python extract_workbook.py model.xlsx -o workbook.json [--max-rows 400] [--sheet "P&L"]

Deterministic half of the kt-corporate-memo pipeline. It does NOT interpret the
workbook; it maps it, so the authoring step argues from real cells instead of guesses.

Emits per sheet: dimensions, detected rectangular tables (header + rows), labeled
scalar figures (a text label with a number beside it), formulas, and number formats.
Every extracted value carries its A1 trace so the memo can cite it.
"""
import argparse, json, re, sys
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

CURRENCY = re.compile(r'[$€£¥]|\bUSD\b')
PERCENT = re.compile(r'%')


def _kind(fmt: str) -> str:
    if not fmt or fmt == "General":
        return "number"
    if PERCENT.search(fmt):
        return "percent"
    if CURRENCY.search(fmt):
        return "currency"
    if re.search(r'[ymd]{2,}', fmt, re.I):
        return "date"
    if fmt.strip().lower().endswith('x"') or fmt.strip().endswith('x'):
        return "multiple"
    return "number"


def _clean(v):
    """JSON-safe scalar."""
    if v is None:
        return None
    if isinstance(v, (int, float, str, bool)):
        return v
    return str(v)


def _is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _row_signature(vals):
    """(n_text, n_num) for a row of values."""
    t = sum(1 for v in vals if isinstance(v, str) and v.strip())
    n = sum(1 for v in vals if _is_num(v))
    return t, n


def detect_tables(grid, max_tables=8, min_rows=2):
    """
    Find rectangular regions that look like tables: a mostly-text header row
    followed by rows carrying numbers in the same columns.

    grid: list of rows, each a list of (value, coord, fmt) or None
    """
    tables, r, nrows = [], 0, len(grid)
    while r < nrows and len(tables) < max_tables:
        row = grid[r]
        vals = [c[0] if c else None for c in row]
        t, n = _row_signature(vals)
        # header candidate: >=2 text cells, no more than one number
        if t >= 2 and n <= 1:
            cols = [i for i, v in enumerate(vals) if v is not None and str(v).strip()]
            body = []
            rr = r + 1
            while rr < nrows:
                nxt = [c[0] if c else None for c in grid[rr]]
                if all(v is None or str(v).strip() == "" for v in nxt):
                    break
                _, nn = _row_signature([nxt[i] for i in cols if i < len(nxt)])
                if nn == 0 and len(body) >= min_rows:
                    break
                body.append(rr)
                rr += 1
            if len(body) >= min_rows:
                header = [str(vals[i]).strip() for i in cols]
                rows, traces = [], []
                for br in body:
                    cells = [grid[br][i] if i < len(grid[br]) else None for i in cols]
                    rows.append([_clean(c[0]) if c else None for c in cells])
                    traces.append([c[1] if c else None for c in cells])
                fmts = []
                for i in cols:
                    sample = next((grid[br][i] for br in body
                                   if i < len(grid[br]) and grid[br][i] and _is_num(grid[br][i][0])), None)
                    fmts.append(_kind(sample[2]) if sample else "text")
                tables.append({
                    "header_row": r + 1,
                    "columns": header,
                    "column_kinds": fmts,
                    "rows": rows,
                    "row_traces": traces,
                    "range": f"{get_column_letter(cols[0]+1)}{r+1}:"
                             f"{get_column_letter(cols[-1]+1)}{body[-1]+1}",
                })
                r = body[-1] + 1
                continue
        r += 1
    return tables


def detect_figures(grid, limit=80):
    """Label-then-number pairs: 'Revenue' | 12400000  ->  a citable scalar."""
    out = []
    for row in grid:
        for i, cell in enumerate(row):
            if not cell or not isinstance(cell[0], str):
                continue
            label = cell[0].strip()
            if not label or len(label) > 60:
                continue
            for j in range(i + 1, min(i + 4, len(row))):
                nxt = row[j]
                if nxt and _is_num(nxt[0]):
                    out.append({
                        "label": label,
                        "value": nxt[0],
                        "trace": nxt[1],
                        "kind": _kind(nxt[2]),
                    })
                    break
            if len(out) >= limit:
                return out
    return out


def extract(path, max_rows=400, max_cols=40, only_sheet=None):
    wb_f = load_workbook(path, data_only=False)   # formulas
    wb_v = load_workbook(path, data_only=True)    # cached values

    doc = {"source_file": path.split("/")[-1], "sheets": [], "named_ranges": [], "warnings": []}

    for name, dest in getattr(wb_f, "defined_names", {}).items() if hasattr(wb_f, "defined_names") else []:
        try:
            doc["named_ranges"].append({"name": name, "refers_to": str(dest.value)})
        except Exception:
            pass

    for ws_f in wb_f.worksheets:
        if only_sheet and ws_f.title != only_sheet:
            continue
        ws_v = wb_v[ws_f.title]
        nrows, ncols = min(ws_f.max_row, max_rows), min(ws_f.max_column, max_cols)
        if ws_f.max_row > max_rows:
            doc["warnings"].append(
                f"'{ws_f.title}' truncated at {max_rows} of {ws_f.max_row} rows")
        if ws_f.max_column > max_cols:
            doc["warnings"].append(
                f"'{ws_f.title}' truncated at {max_cols} of {ws_f.max_column} columns")

        grid, formulas, blank = [], [], True
        for ri in range(1, nrows + 1):
            row = []
            for ci in range(1, ncols + 1):
                cf, cv = ws_f.cell(ri, ci), ws_v.cell(ri, ci)
                val = cv.value if cv.value is not None else cf.value
                coord = f"{ws_f.title}!{get_column_letter(ci)}{ri}"
                if val is None:
                    row.append(None)
                    continue
                blank = False
                row.append((val, coord, cf.number_format or "General"))
                if isinstance(cf.value, str) and cf.value.startswith("="):
                    if len(formulas) < 60:
                        formulas.append({"trace": coord, "formula": cf.value,
                                         "value": _clean(cv.value)})
            grid.append(row)

        if blank:
            continue

        doc["sheets"].append({
            "name": ws_f.title,
            "dimensions": f"{ws_f.max_row} rows x {ws_f.max_column} cols",
            "tables": detect_tables(grid),
            "figures": detect_figures(grid),
            "formulas": formulas,
        })

    if not doc["sheets"]:
        doc["warnings"].append("No non-empty sheets found.")
    return doc


def main():
    ap = argparse.ArgumentParser(description="Extract an .xlsx into a JSON brief.")
    ap.add_argument("xlsx")
    ap.add_argument("-o", "--out", default="workbook.json")
    ap.add_argument("--max-rows", type=int, default=400)
    ap.add_argument("--max-cols", type=int, default=40)
    ap.add_argument("--sheet", default=None, help="limit to one sheet")
    a = ap.parse_args()

    doc = extract(a.xlsx, a.max_rows, a.max_cols, a.sheet)
    with open(a.out, "w") as f:
        json.dump(doc, f, indent=2, default=str)

    print(f"wrote {a.out}")
    for s in doc["sheets"]:
        print(f"  {s['name']}: {len(s['tables'])} tables, "
              f"{len(s['figures'])} figures, {len(s['formulas'])} formulas")
    for w in doc["warnings"]:
        print(f"  ! {w}", file=sys.stderr)


if __name__ == "__main__":
    main()
